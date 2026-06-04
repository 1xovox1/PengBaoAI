import argparse
from dataclasses import dataclass
from pathlib import Path

from bs4 import BeautifulSoup
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


@dataclass(frozen=True)
class SlideSpec:
    title: str
    bullets: list[str]
    speaker_notes: str | None


def _read_html(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def _extract_slides(html: str) -> list[SlideSpec]:
    soup = BeautifulSoup(html, "html.parser")
    slides: list[SlideSpec] = []

    for section in soup.select("section.slide"):
        title_el = section.find(["h2", "h1"])
        title = title_el.get_text(strip=True) if title_el else "Untitled"

        bullets: list[str] = []
        for ul in section.find_all("ul"):
            for li in ul.find_all("li", recursive=False):
                text = li.get_text(" ", strip=True)
                if text:
                    bullets.append(text)

        speaker = None
        speaker_el = section.select_one(".speaker")
        if speaker_el:
            speaker = speaker_el.get_text(" ", strip=True)

        slides.append(SlideSpec(title=title, bullets=bullets, speaker_notes=speaker))

    if not slides:
        raise ValueError("No slides found. Make sure HTML contains <section class='slide'> blocks.")

    return slides


def _set_font(run, *, name: str, size_pt: int, color_rgb: tuple[int, int, int] | None = None, bold: bool = False):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if color_rgb is not None:
        run.font.color.rgb = RGBColor(*color_rgb)


def build_pptx(slides: list[SlideSpec], out_path: Path) -> None:
    prs = Presentation()

    # Use built-in layouts:
    # 0 = title slide, 1 = title and content
    title_layout = prs.slide_layouts[0]
    content_layout = prs.slide_layouts[1]

    # Prefer a common Chinese font; fallback handled by Office if unavailable.
    font_name = "Microsoft YaHei"

    # Slide 1 as title slide
    first = slides[0]
    s0 = prs.slides.add_slide(title_layout)
    title_shape = s0.shapes.title
    subtitle_shape = s0.placeholders[1] if len(s0.placeholders) > 1 else None

    title_shape.text = first.title
    # Style title
    if title_shape.text_frame.paragraphs and title_shape.text_frame.paragraphs[0].runs:
        _set_font(title_shape.text_frame.paragraphs[0].runs[0], name=font_name, size_pt=34, bold=True)

    if subtitle_shape is not None:
        subtitle_shape.text = "战略版（由 HTML 自动生成）"
        p = subtitle_shape.text_frame.paragraphs[0]
        if p.runs:
            _set_font(p.runs[0], name=font_name, size_pt=16, color_rgb=(90, 90, 90))
        p.alignment = PP_ALIGN.LEFT

    if first.speaker_notes:
        s0.notes_slide.notes_text_frame.text = first.speaker_notes

    # Remaining slides
    for spec in slides[1:]:
        slide = prs.slides.add_slide(content_layout)

        # Title
        t = slide.shapes.title
        t.text = spec.title
        tp = t.text_frame.paragraphs[0]
        if tp.runs:
            _set_font(tp.runs[0], name=font_name, size_pt=28, bold=True)

        # Body
        body = slide.shapes.placeholders[1].text_frame
        body.clear()
        body.word_wrap = True

        for i, b in enumerate(spec.bullets):
            p = body.paragraphs[0] if i == 0 else body.add_paragraph()
            p.text = b
            p.level = 0
            p.space_after = Pt(6)
            if p.runs:
                _set_font(p.runs[0], name=font_name, size_pt=16)

        if spec.speaker_notes:
            slide.notes_slide.notes_text_frame.text = spec.speaker_notes

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)


def main() -> int:
    parser = argparse.ArgumentParser(description="Convert slide-structured HTML into PPTX.")
    parser.add_argument("--html", required=True, help="Path to HTML file with <section class='slide'> blocks.")
    parser.add_argument("--out", required=True, help="Output .pptx path.")
    args = parser.parse_args()

    html_path = Path(args.html).expanduser().resolve()
    out_path = Path(args.out).expanduser().resolve()

    html = _read_html(html_path)
    slides = _extract_slides(html)
    build_pptx(slides, out_path)

    print(f"OK: wrote {out_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

